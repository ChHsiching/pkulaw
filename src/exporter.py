import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def _sanitize(value):
    if isinstance(value, str):
        return _ILLEGAL_RE.sub("", value)
    return value


BASE_COLUMNS = [
    "gid",
    "url",
    "title",
    "案例层级",
    "法宝引证码",
    "时效性",
    "案由",
    "案号",
    "文书类型",
    "公开类型",
    "审理法院",
    "审结日期",
    "案件类型",
    "审理程序",
    "案例发文",
    "案例编号",
    "发布日期",
    "来源",
    "刑罚",
    "指控罪名",
    "判定罪名",
    "审理法官",
    "代理律师/律所",
    "权责关键词",
]

WIDE_COLUMNS = {"gid", "full_text", "url", "刑罚"}

METADATA_KEYS = set(BASE_COLUMNS) | {"法宝引证码", "时效性", "案例层级"}


def _build_column_order(cases: list[dict]) -> list[str]:
    section_labels = set()
    for case in cases:
        for key in case:
            if (
                key not in METADATA_KEYS
                and key != "full_text"
                and key not in {"gid", "url", "title"}
            ):
                if case.get(key, ""):
                    section_labels.add(key)

    base = [c for c in BASE_COLUMNS if c not in {"法宝引证码", "时效性"}]
    existing_base = []
    for c in base:
        if c in {"案例层级"}:
            continue
        existing_base.append(c)

    ordered = ["gid", "url", "title", "案例层级", "法宝引证码", "时效性"]
    ordered += [c for c in BASE_COLUMNS if c not in ordered]
    ordered += sorted(section_labels - set(ordered))
    ordered.append("full_text")
    return ordered


def export_json(cases: list[dict], path: Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(cases, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def export_excel(cases: list[dict], path: Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    columns = _build_column_order(cases)

    wb = Workbook()
    ws = wb.active
    ws.title = "司法案例"

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True)

    for row_idx, case in enumerate(cases, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = _sanitize(case.get(col_name, ""))
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, col_name in enumerate(columns, 1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        if col_name in WIDE_COLUMNS or col_name == "full_text":
            ws.column_dimensions[letter].width = 60
        else:
            ws.column_dimensions[letter].width = 20

    wb.save(str(path))
